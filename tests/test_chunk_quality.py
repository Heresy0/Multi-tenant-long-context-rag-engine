"""Regressions for conservative cleaning, role hints and complete rule units."""
import json
import unicodedata
from pathlib import Path

import pytest
from docx import Document as WordDocument
from langchain_core.documents import Document
from langchain_text_splitters import RecursiveCharacterTextSplitter

from backend.app.documents.chunk_quality import (
    CLEANING_VERSION, annotate_chunks, chunk_body, classify_chunk, summarize_quality,
)
from backend.app.documents.document_splitter import (
    ContentUnit, split_content_units, split_rule_text, table_rows,
)
from backend.app.documents.pdf_splitter import (
    ExtractedPage, collect_pdf_units, normalize_pdf_text, remove_repeated_marginalia, split_pdf,
)
from backend.app.knowledge.context_builder import ContextBuilder
from backend.app.knowledge.evidence_quality import evidence_role
from backend.app.knowledge.rag import split_file


@pytest.mark.parametrize("page", ["1", "1-2"])
@pytest.mark.parametrize("body", [
    "星海科技有限公司 | 内部资料", "星海科技有限公司", "文档编号：XH-PMO-2026-14；状态：现行",
])
def test_generated_pdf_wrapper_does_not_turn_cover_into_body(page, body):
    text = f"文档：验收纪要\n章节：文档说明\n页码：{page}\n\n{body}"
    doc = Document(page_content=text, metadata={"document_name": "验收纪要", "section_path": ""})
    assert classify_chunk(text, doc.metadata) == "front_matter"
    assert evidence_role(doc) == "front_matter"
    assert chunk_body(text) == unicodedata.normalize("NFKC", body)


@pytest.mark.parametrize("text", [
    "文档：审批负责人\n章节：业务规则\n页码：1\n必须先审批。",  # No generated blank separator.
    "星海科技有限公司 | 内部资料\n客户已于9月29日签署最终验收。",
    "文档编号：RULE-001；必须先取得审批。",
    "星海科技有限公司负责审批。",
    "未标章节，但401不得重试。",
])
def test_unknown_or_mixed_business_prose_is_not_removed(text):
    assert classify_chunk(text, {"section_path": ""}) == "content"


def test_stored_front_matter_hint_never_overrides_actual_rule():
    doc = Document(page_content="401不得自动重试。", metadata={"evidence_role": "front_matter"})
    assert evidence_role(doc) == "content"


def test_code_that_looks_like_metadata_is_still_content():
    assert classify_chunk("版本: 1.0", {"chunk_type": "code"}) == "content"


def test_ocr_marks_review_without_correcting_or_certifying_name():
    chunks = annotate_chunks([
        Document(page_content="验证负责人：周末", metadata={"parser": "ocr", "page_start": 1}),
        Document(page_content="正文。", metadata={"parser": "pypdf"}),
        Document(page_content="另一页OCR", metadata={"parser": "pypdf+ocr"}),
    ])
    assert chunks[0].page_content == "验证负责人：周末"
    assert chunks[0].metadata["quality"] == {"review_required": True, "review_reasons": ["ocr_unverified"]}
    assert chunks[1].metadata["quality"]["review_required"] is False
    summary = summarize_quality(chunks)
    assert summary["ocr_chunks"] == 2
    assert summary["review_required"] is True
    assert summary["cleaning_version"] == CLEANING_VERSION


@pytest.mark.parametrize("text", ["650", "550", "2025", "-10", "09:12", "1.4"])
def test_pdf_normalization_preserves_business_numbers(text):
    assert normalize_pdf_text(f"限额或时间\n{text}\n其他规则") == f"限额或时间\n{text}\n其他规则"


@pytest.mark.parametrize("rule", [
    "普通员工住宿上限650元。", "版本：1.4", "不得进入正式验收。", "审批人|时限",
    "姓名林澈。", "维护负责人林澈", "如果失败需要核查", "此项不适用。",
    "不自动授予权限", "客户已确认", "只允许沙箱访问", "导出任务在沙箱执行", "联系人林澈",
])
def test_repeated_edge_business_content_survives(rule):
    pages = [ExtractedPage(i, f"{rule}\n中间正文\n第 {i} 页", "pypdf") for i in (1, 2)]
    cleaned = remove_repeated_marginalia(pages)
    assert all(rule in page.text for page in cleaned)
    assert all(f"第 {page.number} 页" not in page.text for page in cleaned)


def test_numeric_normalization_does_not_erase_different_historical_limits():
    pages = [ExtractedPage(1, "住宿限额550元\n正文", "pypdf"),
             ExtractedPage(2, "住宿限额650元\n正文", "pypdf")]
    cleaned = remove_repeated_marginalia(pages)
    assert "550元" in cleaned[0].text and "650元" in cleaned[1].text


def test_bare_numbers_and_body_page_references_are_not_assumed_to_be_footers():
    text = "第 1 页\n规则见上页\n650\n第 1 页"
    cleaned = remove_repeated_marginalia([ExtractedPage(1, text, "pypdf")])[0].text
    assert cleaned == "第 1 页\n规则见上页\n650"


def test_pdf_cover_is_not_coalesced_with_unheaded_rule():
    units = collect_pdf_units([ExtractedPage(1, "星海科技有限公司 | 内部资料\n\n401不得重试。", "pypdf")])
    assert len(units) == 2
    assert units[1].content == "401不得重试。"


@pytest.mark.parametrize("suffix", [".txt", ".md"])
def test_new_text_chunks_carry_quality_metadata(tmp_path, suffix):
    path = tmp_path / f"规范{suffix}"
    path.write_text("星海科技有限公司\n\n401不得重试。", encoding="utf-8")
    chunks = split_file(path)
    assert [doc.metadata["evidence_role"] for doc in chunks] == ["front_matter", "content"]
    assert all(doc.metadata["cleaning_version"] == CLEANING_VERSION for doc in chunks)
    assert all(not doc.metadata["quality"]["review_required"] for doc in chunks)


def test_ocr_fallback_propagates_page_and_review_marker():
    chunks = split_pdf(Path("tests/fixtures/pdfs/scanned_notice.pdf"),
                       ocr_page=lambda *_: "1 恢复结果\n维护负责人林澈，验证负责人周末。")
    assert chunks[0].metadata["quality"]["review_required"] is True
    assert chunks[0].metadata["page_start"] == 1
    assert "周末" in chunks[0].page_content


def test_short_conditional_rule_is_not_split_away_from_exception():
    text = "仅当满足幂等条件时允许自动重试。" + "规则说明。" * 100 + "例外：401不得自动重试。"
    splitter = RecursiveCharacterTextSplitter(chunk_size=450, chunk_overlap=60)
    assert 450 < len(text) <= 900
    assert split_rule_text(text, splitter) == [text]


@pytest.mark.parametrize("suffix", [".md", ".txt"])
def test_text_condition_list_keeps_introduction_and_all_conditions(tmp_path, suffix):
    path = tmp_path / f"验收{suffix}"
    heading = "# 验收" if suffix == ".md" else "1 验收"
    path.write_text(f"{heading}\n必须同时满足以下条件：\n- 存在可接受替代方案\n- 不阻断核心流程\n"
                    "- 客户书面接受\n- 明确修复日期\n\n其他事项。", encoding="utf-8")
    chunks = split_file(path)
    assert chunks[0].metadata["chunk_type"] == "rule_group"
    assert all(text in chunks[0].page_content for text in ["必须同时满足", "替代方案", "核心流程", "书面接受", "修复日期"])
    assert "其他事项" not in chunks[0].page_content


def test_docx_condition_list_keeps_introduction_and_all_conditions(tmp_path):
    word = WordDocument()
    word.add_heading("验收", level=1)
    word.add_paragraph("必须同时满足以下条件：")
    for text in ["存在可接受替代方案", "不阻断核心流程", "客户书面接受", "明确修复日期"]:
        word.add_paragraph(text, style="List Bullet")
    path = tmp_path / "验收.docx"
    word.save(path)
    chunks = split_file(path)
    assert len(chunks) == 1
    assert chunks[0].metadata["chunk_type"] == "rule_group"
    assert chunks[0].metadata["parser"] == "python-docx"
    assert all(text in chunks[0].page_content for text in ["替代方案", "核心流程", "书面接受", "修复日期"])


def test_long_condition_group_repeats_intro_and_retains_shared_parent(tmp_path):
    intro = "必须同时满足以下条件："
    units = [ContentUnit("paragraph", intro, ("验收",)),
             *[ContentUnit("list", f"- 条件{i}：" + "说明" * 60, ("验收",)) for i in range(15)]]
    chunks = split_content_units(units, path=tmp_path / "规则.md", document_name="规则")
    assert len(chunks) > 1
    assert len({chunk.metadata["parent_key"] for chunk in chunks}) == 1
    assert all(chunk_body(chunk.page_content).startswith(intro.replace("：", ":")) for chunk in chunks)
    assert all(len(chunk_body(chunk.page_content)) <= 900 for chunk in chunks)
    assert "条件14" in chunks[-1].page_content


def test_long_rule_intro_counts_toward_rule_chunk_budget():
    intro = "必须满足以下条件" + "说明" * 80 + "："
    text = intro + "\n" + "条件详情。" * 500
    splitter = RecursiveCharacterTextSplitter(chunk_size=2000, chunk_overlap=100)
    parts = split_rule_text(text, splitter, rule_intro=intro)
    assert len(parts) > 1
    assert all(part.startswith(intro) and len(part) <= 900 for part in parts)


def test_condition_binding_does_not_cross_heading_or_table(tmp_path):
    units = [ContentUnit("paragraph", "必须满足以下条件：", ("A",)),
             ContentUnit("list", "- 另一章节的规则", ("B",)),
             ContentUnit("table", ["等级|条件", "D3|客户书面接受"], ("B",))]
    chunks = split_content_units(units, path=tmp_path / "规则.docx", document_name="规则")
    assert len(chunks) == 3
    assert all(chunk.metadata["chunk_type"] != "rule_group" for chunk in chunks)


def test_pdf_condition_list_keeps_page_range_and_applicability():
    pages = [ExtractedPage(1, "1 验收\n必须同时满足以下条件：\n- 存在替代方案", "pypdf"),
             ExtractedPage(2, "- 不阻断流程\n- 客户书面接受", "ocr")]
    units = collect_pdf_units(pages)
    assert len(units) == 1
    assert units[0].kind == "rule_group"
    assert (units[0].page_start, units[0].page_end) == (1, 2)
    assert set(units[0].parsers) == {"ocr", "pypdf"}
    assert "必须同时满足以下条件" in str(units[0].content)


def test_docx_table_preserves_semantic_spaces():
    word = WordDocument()
    table = word.add_table(rows=2, cols=2)
    table.cell(0, 0).text, table.cell(0, 1).text = "字段", "示例"
    table.cell(1, 0).text, table.cell(1, 1).text = "认证", "Bearer access_token\n第二行"
    assert table_rows(table)[1] == "认证|Bearer access_token 第二行"


def test_observed_hailan_trace_keeps_actual_acceptance_after_cover_demotion():
    path = Path("evals/reports/system/20261004T173241439459Z/traces.jsonl")
    if not path.exists():
        pytest.skip("Local sensitive trace is not committed")
    trace = next(json.loads(line) for line in path.read_text(encoding="utf-8").splitlines()
                 if line.strip() and json.loads(line).get("id") == "XH-031")
    documents = [Document(id=row["chunk_id"], page_content=row["content"], metadata={
        "source_id": row["document_name"], "source": row["document_name"],
        "document_name": row["document_name"], "section_path": row.get("section_path"),
    }) for row in trace["stages"]["rerank"]]
    context = ContextBuilder().build(documents, question=trace["query"])
    assert "客户于2026年9月29日签署最终验收" in context.text
    assert "客户要求新增跨部门审批代理功能" in context.text
    assert all("星海科技有限公司 | 内部资料" != chunk_body(item.content) for item in context.items)
    assert len([item for item in context.items if item.document_name.startswith("13_")]) <= 4
    assert context.total_characters <= 12000


def test_synthetic_pdf_cover_does_not_displace_actual_acceptance_without_local_trace():
    documents = [Document(id=str(i), page_content=f"文档：验收纪要\n章节：{section or '文档说明'}\n页码：1\n\n{body}",
                         metadata={"source_id": "acceptance", "document_name": "验收纪要", "section_path": section})
                 for i, (section, body) in enumerate([
                     ("2 范围", "原计划9月25日，调整到9月29日。"), ("", "星海科技有限公司 | 内部资料"),
                     ("3 记录", "审批测试通过。"), ("1 会议", "会议日期9月29日。"),
                     ("4 决议", "客户于9月29日签署最终验收。")])]
    context = ContextBuilder().build(documents, question="最后何时验收？")
    assert "签署最终验收" in context.text
    assert [item.chunk_id for item in context.items] == ["0", "2", "3", "4"]


@pytest.mark.parametrize("show_content", [False, True])
def test_inspection_cli_is_read_only_and_hides_content_by_default(tmp_path, monkeypatch, capsys, show_content):
    from scripts.inspect_document_chunks import main
    path = tmp_path / "规范.txt"
    path.write_text("秘密业务规则：401不得重试。", encoding="utf-8")
    before = path.read_bytes()
    argv = ["inspect_document_chunks", str(path)] + (["--show-content"] if show_content else [])
    monkeypatch.setattr("sys.argv", argv)
    monkeypatch.setattr("backend.app.knowledge.rag.create_embeddings",
                        lambda *_args, **_kwargs: pytest.fail("Inspection must not create embeddings"))
    assert main() == 0
    result = json.loads(capsys.readouterr().out)
    assert result["chunking_version"] == "structured-v4"
    assert result["chunk_count"] == 1
    assert ("content" in result["chunks"][0]) == show_content
    assert path.read_bytes() == before
    assert list(tmp_path.iterdir()) == [path]
