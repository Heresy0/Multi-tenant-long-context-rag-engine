"""Strip corpus provenance from business bodies, preserving the original layout.

Offline only. Creates a complete backup before changing the active pack.
Run using the bundled document runtime with PyMuPDF available. DOCX layout
requires the user's explicitly accepted manual Word review in this environment.
This is a corpus-specific allowlist, NOT a generic ingestion sanitizer.
"""
from __future__ import annotations

import argparse
import csv
import hashlib
import json
import re
import shutil
import zipfile
from pathlib import Path

from docx import Document
from pypdf import PdfReader
import fitz

ROOT = Path(__file__).resolve().parents[1]
CORPUS = ROOT / 'output/pdf/xinghai-department-corpus-v3'
NOTICE = '虚构测试资料，仅用于知识库开发与评估，不具有实际管理或法律效力。'
# Keep sandbox-data constraints and API/example credentials warnings: these are
# useful business safeguards, not a corpus-wide disclaimer.
REMOVALS = (
    NOTICE,
    '文件说明：模拟企业文件，仅用于检索增强生成系统的开发、测试与评估，不具有实际管理或法律效力。',
    '企业 RAG 测试语料',
    '海岚制造有限公司及人员均为虚构。',
    '评测语料只保留虚构记录。',
    '客户名称、工单号和人员均为虚构，案例不得用于推断真实客户的事件。',
    '，本文只是虚构内部操作测试资料',
    '虚构单号',
    '虚构的',
    '虚构测试资料',
    '虚构',
    '模拟规模',
)
REPLACEMENTS = {'虚构单号': '单号', '模拟规模': '规模'}


def digest(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def compact(text):
    return re.sub(r'\s+', '', text)


def clean_text(text, *, keep_example_warning=False):
    for value in REMOVALS:
        if keep_example_warning and value == '虚构':
            # Local API introductory paragraph: keep the exact sample warning.
            continue
        text = text.replace(value, REPLACEMENTS.get(value, ''))
    return text


def docx_text(path):
    doc = Document(path)
    return '\n'.join(''.join(node.text or '' for node in p.xpath('.//w:t'))
                     for p in doc.element.xpath('.//w:p'))


def edit_docx(path):
    doc = Document(path)
    for paragraph in doc.paragraphs:
        if '示例中的地址、密钥和业务数据' in paragraph.text:
            continue
        expected = clean_text(paragraph.text)
        if expected == paragraph.text:
            continue
        if not expected.strip():
            paragraph._p.getparent().remove(paragraph._p)
        else:
            # These known notices each occupy one run. Refuse a broad formatting
            # rewrite when a future source differs from the reviewed fixtures.
            assert len(paragraph.runs) == 1, paragraph.text
            paragraph.runs[0].text = expected
    doc.save(path)


def edit_pdf(path):
    doc = fitz.open(path)
    changes = []
    for page in doc:
        reflow = []
        if not page.get_text().strip():
            # Known scan banner only; retain all other image content. Coordinates
            # map the original 1654x2339 scan to A4, including rotation tolerance.
            assert path.name == '12_生产维护通知_扫描件.pdf' and len(doc) == 1
            rect = fitz.Rect(38, 120, page.rect.width - 38, 157)
            page.add_redact_annot(rect, fill=(250/255, 250/255, 247/255))
            changes.append(dict(page=page.number + 1, text='扫描件测试说明', rect=list(rect)))
        else:
            # Remove characters rather than paint overlays; the old notice must
            # not remain searchable/extractable under a white rectangle.
            terms = [v for v in REMOVALS if v not in REPLACEMENTS]
            terms += ['模拟', '虚构测试资料']
            chars = []
            blocks = []
            for block in page.get_text('rawdict')['blocks']:
                if not block.get('lines'):
                    continue
                # Some source PDFs expose each Chinese line as its own block.
                # Join adjacent 11pt body lines, never headings or table cells.
                size = block['lines'][0]['spans'][0]['size']
                if (blocks and abs(size - 11) < .1 and
                    abs(blocks[-1]['lines'][0]['spans'][0]['size'] - size) < .1 and
                    abs(block['bbox'][0] - blocks[-1]['bbox'][0]) < .5 and
                    0 <= block['bbox'][1] - blocks[-1]['bbox'][3] < 9):
                    blocks[-1]['lines'].extend(block['lines'])
                    blocks[-1]['bbox'] = tuple(fitz.Rect(blocks[-1]['bbox']) | fitz.Rect(block['bbox']))
                else:
                    blocks.append(block)
            for block in blocks:
                block_chars = [c for line in block.get('lines', []) for span in line['spans'] for c in span['chars']]
                block_text = ''.join(c['c'] for c in block_chars)
                if any(v in block_text for v in ('虚构企业', '模拟规模', '虚构的海岚', '虚构事件',
                                                '本文只是虚构', '客户名称、工单号和人员均为')):
                    rect = fitz.Rect(block['bbox'])
                    page.add_redact_annot(rect, fill=(1, 1, 1))
                    reflow.append((rect, clean_text(block_text), block['lines'][0]['spans'][0]['size']))
                    changes.append(dict(page=page.number+1, text='业务段落局部重新排版'))
                    continue
                for line in block.get('lines', []):
                    for span in line['spans']:
                        chars.extend(c for c in span['chars'] if not c['c'].isspace())
            flat = ''.join(c['c'] for c in chars)
            for term in terms:
                needle = compact(term)
                offset = 0
                while (start := flat.find(needle, offset)) >= 0:
                    # Per-character rectangles also handle a sentence wrapped
                    # across Chinese lines with no inserted whitespace.
                    rects = [fitz.Rect(c['bbox']) for c in chars[start:start + len(needle)]]
                    for rect in rects:
                        rect.x0 += 0.1
                        rect.x1 -= 0.1
                        rect.y0 += 1
                        rect.y1 -= 1
                        page.add_redact_annot(rect, fill=(1, 1, 1))
                    changes.append(dict(page=page.number + 1, text=term, characters=len(rects)))
                    offset = start + len(needle)
        page.apply_redactions(images=2, graphics=0, text=0)
        for rect, text, size in reflow:
            # Keep the same body font and page location; reflow only the edited
            # paragraphs so deleted prefixes do not leave conspicuous holes.
            rect.y1 += 10
            remaining = page.insert_textbox(rect, text, fontsize=size, fontname='businessbody',
                fontfile='C:/Windows/Fonts/msyh.ttc', lineheight=1.45)
            assert remaining >= 0, (path, text, remaining)
    assert changes, path
    if reflow:
        doc.subset_fonts(fallback=True)
    # PyMuPDF's TTC embedding can encode astral Unicode destinations with odd
    # hex lengths. Convert these to valid UTF-16BE rather than leave a malformed
    # ToUnicode map for downstream PDF readers.
    for page in doc:
        for font in page.get_fonts():
            kind, reference = doc.xref_get_key(font[0], 'ToUnicode')
            if kind == 'xref':
                xref = int(reference.split()[0])
                stream = doc.xref_stream(xref)
                corrected = re.sub(rb'<([0-9a-fA-F]{5,6})>',
                    lambda m: b'<' + chr(int(m[1], 16)).encode('utf-16-be').hex().encode() + b'>', stream)
                if corrected != stream:
                    doc.update_stream(xref, corrected)
    target = path.with_suffix('.cleaning.pdf')
    doc.save(target, garbage=4, deflate=True)
    doc.close()
    target.replace(path)
    return changes


def run(backup):
    global CORPUS
    assert backup.is_relative_to(ROOT / 'tmp'), 'Backup must remain within project tmp'
    manifest = json.loads((CORPUS / 'manifest.json').read_text(encoding='utf-8'))
    if manifest.get('body_provenance_policy') == 'central-readme-only-2026-10-05':
        print('Active corpus is already cleaned; no files changed.')
        return
    assert manifest['total_documents'] == 30
    for row in manifest['documents']:
        assert digest(CORPUS / row['file']) == row['sha256'], row['file']
    if not backup.exists():
        shutil.copytree(CORPUS, backup / 'corpus')
    else:
        for row in manifest['documents']:
            assert digest(backup / 'corpus' / row['file']) == row['sha256']
    legacy = ROOT / 'sample_docs/fictional_enterprise/06_开放平台API接入与故障处理技术手册.docx'
    if not (backup / legacy.name).exists():
        shutil.copy2(legacy, backup / legacy.name)
    assert digest(legacy) == digest(backup / legacy.name)
    active = CORPUS
    CORPUS = backup.parent / 'candidate-corpus'
    shutil.copytree(backup / 'corpus', CORPUS, dirs_exist_ok=True)
    legacy_original = legacy
    legacy = backup.parent / 'candidate-legacy' / legacy.name
    legacy.parent.mkdir(exist_ok=True)
    shutil.copy2(legacy_original, legacy)
    records = []
    for path in [CORPUS / row['file'] for row in manifest['documents']] + [legacy]:
        old_hash = digest(path)
        if path.suffix == '.pdf':
            before = '\n'.join(p.extract_text() or '' for p in PdfReader(path).pages)
            with fitz.open(path) as pdf:
                ordered_before = '\n'.join(p.get_text(sort=True) for p in pdf)
            removals = edit_pdf(path)
            after = '\n'.join(p.extract_text() or '' for p in PdfReader(path).pages)
            if before.strip():
                # Removing text cannot change any other business character.
                expected = compact(ordered_before)
                for term in REMOVALS:
                    expected = expected.replace(compact(term), compact(REPLACEMENTS.get(term, '')))
                with fitz.open(path) as pdf:
                    ordered_after = '\n'.join(p.get_text(sort=True) for p in pdf)
                assert expected == compact(ordered_after), (path, expected, compact(ordered_after))
        elif path.suffix == '.docx':
            before = docx_text(path)
            edit_docx(path)
            after = docx_text(path)
            expected = '\n'.join(clean_text(line, keep_example_warning='示例中的地址、密钥和业务数据' in line)
                                 for line in before.splitlines())
            assert compact(expected) == compact(after), path
            removals = []
        else:
            before = path.read_text(encoding='utf-8')
            after = '\n'.join(clean_text(line, keep_example_warning='示例中的地址、密钥和业务数据' in line)
                               for line in before.splitlines()) + '\n'
            after = re.sub(r'\n{3,}', '\n\n', after)
            path.write_text(after, encoding='utf-8')
            removals = []
        logical = legacy_original if path == legacy else active / path.relative_to(CORPUS)
        records.append(dict(file=logical.relative_to(ROOT).as_posix(), old_sha256=old_hash,
                            sha256=digest(path), bytes=path.stat().st_size,
                            content_characters=len(after), pdf_removals=removals))
    by_file = {r['file']: r for r in records}
    for row in manifest['documents']:
        record = by_file[(active / row['file']).relative_to(ROOT).as_posix()]
        row.update(sha256=record['sha256'], bytes=record['bytes'])
        row['text_characters'] = record['content_characters']
        row['source_characters'] = record['content_characters']
        if row['format'] == 'pdf':
            row['embedded_text_characters'] = record['content_characters']
        row['body_provenance_removed'] = True
    truth_path = CORPUS / 'evaluation/scan_ground_truth.json'
    truth = json.loads(truth_path.read_text(encoding='utf-8'))
    truth['text'] = truth['text'].replace('虚构测试资料\n', '')
    truth_path.write_text(json.dumps(truth, ensure_ascii=False, indent=2) + '\n', encoding='utf-8')
    scan = next(r for r in manifest['documents'] if Path(r['file']).name == truth['document'])
    scan['source_characters'] = scan['text_characters'] = len(truth['text'])
    manifest['source_characters'] = sum(r['source_characters'] for r in manifest['documents'])
    manifest['body_provenance_policy'] = 'central-readme-only-2026-10-05'
    (CORPUS / 'manifest.json').write_text(json.dumps(manifest, ensure_ascii=False, indent=2) + '\n', encoding='utf-8')
    for file in CORPUS.glob('*.csv'):
        with file.open(encoding='utf-8-sig', newline='') as stream:
            reader = csv.DictReader(stream)
            fields, rows = reader.fieldnames, list(reader)
        for row in rows:
            relative = row.get('file') or row.get('相对路径')
            if relative:
                original = next(r for r in manifest['documents'] if r['file'] == relative)
                for key in ('sha256', 'bytes'):
                    if key in row:
                        row[key] = original[key]
                if 'SHA256' in row:
                    row['SHA256'] = original['sha256']
        with file.open('w', encoding='utf-8-sig', newline='') as stream:
            writer = csv.DictWriter(stream, fieldnames=fields)
            writer.writeheader()
            writer.writerows(rows)
    # Preserve all reference answers; only normalize the old organizational
    # evidence phrase whose "模拟" prefix was removed from the business body.
    questions = CORPUS / 'evaluation/questions.jsonl'
    questions.write_text(questions.read_text(encoding='utf-8').replace('模拟规模180人', '规模180人'), encoding='utf-8')
    readme = CORPUS / 'README.md'
    text = readme.read_text(encoding='utf-8')
    text += '\n## 正文说明清理 2026年10月5日\n\n公司、人员、客户、事件和业务数值均为合成设定，不具有实际管理或法律效力。该声明集中保留于本说明及manifest，不作为业务正文入库。正文仅删除语料测试说明，业务规则、编号、日期、归档状态和合同模板限制保持不变。API示例地址与密钥提醒保留，避免误用示例配置。DOCX采用用户同意的人工Word排版验收；未声称通过自动分页渲染。\n'
    readme.write_text(text, encoding='utf-8')
    (backup / 'changes.json').write_text(json.dumps(records, ensure_ascii=False, indent=2) + '\n', encoding='utf-8')
    print(json.dumps(dict(edited_files=len(records), pdfs=12, docx=7, backup=str(backup)), ensure_ascii=False))


def promote(backup):
    """Promote only the exact reviewed candidate; refuse changed originals."""
    assert backup.is_relative_to(ROOT / 'tmp')
    records = json.loads((backup / 'changes.json').read_text(encoding='utf-8'))
    replacements = []
    for row in records:
        target = ROOT / row['file'].replace('\\', '/')
        assert target.resolve().is_relative_to(ROOT)
        relative = target.relative_to(CORPUS) if target.is_relative_to(CORPUS) else None
        candidate = backup.parent / 'candidate-corpus' / relative if relative else backup.parent / 'candidate-legacy' / target.name
        assert digest(target) in {row['old_sha256'], row.get('prior_candidate_sha256')} and digest(candidate) == row['sha256']
        replacements.append((candidate, target))
    for candidate, target in replacements:
        shutil.copy2(candidate, target)
    for relative in ('manifest.json', 'README.md', '入库清单.csv',
                     'evaluation/scan_ground_truth.json', 'evaluation/questions.jsonl'):
        shutil.copy2(backup.parent / 'candidate-corpus' / relative, CORPUS / relative)
    archive = ROOT / 'output/pdf/星海科技有限公司_部门知识库30份_v3.zip'
    if archive.exists() and not (backup / archive.name).exists():
        shutil.copy2(archive, backup / archive.name)
    with zipfile.ZipFile(archive, 'w', compression=zipfile.ZIP_DEFLATED) as bundle:
        for path in CORPUS.rglob('*'):
            if path.is_file():
                bundle.write(path, path.relative_to(CORPUS))
    print(json.dumps(dict(promoted=len(replacements), archive=str(archive)), ensure_ascii=False))


def compress_reviewed(backup):
    records = json.loads((backup / 'changes.json').read_text(encoding='utf-8'))
    candidate_root = backup.parent / 'candidate-corpus'
    manifest = json.loads((candidate_root / 'manifest.json').read_text(encoding='utf-8'))
    for row in records:
        path = ROOT / row['file'].replace('\\', '/')
        row.setdefault('prior_candidate_sha256', row['sha256'])
        if path.suffix != '.pdf' or path.name[:2] not in {'07', '13', '25', '27'}:
            continue
        candidate = candidate_root / path.relative_to(CORPUS)
        with fitz.open(candidate) as doc:
            before = '\n'.join(p.get_text(sort=True) for p in doc)
            doc.subset_fonts(fallback=True)
            temporary = candidate.with_suffix('.subset.pdf')
            doc.save(temporary, garbage=4, deflate=True)
        with fitz.open(temporary) as doc:
            assert before == '\n'.join(p.get_text(sort=True) for p in doc)
        temporary.replace(candidate)
        row.update(sha256=digest(candidate), bytes=candidate.stat().st_size)
        entry = next(r for r in manifest['documents'] if Path(r['file']).name == candidate.name)
        entry.update(sha256=row['sha256'], bytes=row['bytes'])
    (backup / 'changes.json').write_text(json.dumps(records, ensure_ascii=False, indent=2) + '\n', encoding='utf-8')
    (candidate_root / 'manifest.json').write_text(json.dumps(manifest, ensure_ascii=False, indent=2) + '\n', encoding='utf-8')
    # Keep the upload list consistent with the newly subsetted PDF hashes.
    path = candidate_root / '入库清单.csv'
    with path.open(encoding='utf-8-sig', newline='') as stream:
        reader = csv.DictReader(stream)
        fields, rows = reader.fieldnames, list(reader)
    for row in rows:
        entry = next(r for r in manifest['documents'] if r['file'] == (row.get('file') or row['相对路径']))
        for key in ('sha256', 'bytes'):
            if key in row:
                row[key] = entry[key]
        if 'SHA256' in row:
            row['SHA256'] = entry['sha256']
    with path.open('w', encoding='utf-8-sig', newline='') as stream:
        writer = csv.DictWriter(stream, fieldnames=fields)
        writer.writeheader()
        writer.writerows(rows)
    print('Subsetted the four edited paragraph fonts; business text unchanged.')


if __name__ == '__main__':
    parser = argparse.ArgumentParser()
    parser.add_argument('--backup', type=Path, required=True)
    parser.add_argument('--promote-reviewed', action='store_true')
    parser.add_argument('--compress-reviewed-fonts', action='store_true')
    args = parser.parse_args()
    if args.compress_reviewed_fonts:
        compress_reviewed(args.backup.resolve())
    elif args.promote_reviewed:
        promote(args.backup.resolve())
    else:
        run(args.backup.resolve())
