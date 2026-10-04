"""Create a separate 30-file edition with six native DOCX replacements.

Requires explicit acceptance of manual Word layout review. Never uploads files,
calls models, modifies the database, or overwrites a previous corpus edition.
Run authoring with the bundled document runtime; use the project validator
separately before packaging.
"""
from __future__ import annotations

import argparse
import csv
import hashlib
import json
import re
import shutil
import zipfile
from collections import Counter
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import RGBColor
from pypdf import PdfReader

from generate_xinghai_corpus import COMPANY, NOTICE, ROOT, dump_json, iter_blocks, make_word
from generate_xinghai_department_corpus import NEW_DOCUMENTS, metadata, word_to_sections

SOURCE = ROOT / "output/pdf/xinghai-department-corpus-v2"
OUTPUT = ROOT / "output/pdf/xinghai-department-corpus-v3"
REPLACE_NUMBERS = {1, 2, 4, 5, 21, 29}
EXPECTED_COUNTS = {"docx": 6, "pdf": 12, "md": 8, "txt": 4}


def sha256(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def normalized(text):
    return "".join(text.split()).replace("|", "")


def word_text(path):
    values = []
    for block in iter_blocks(Document(path)):
        if hasattr(block, "rows"):
            values.extend(" | ".join(cell.text for cell in row.cells) for row in block.rows)
        else:
            values.append(block.text)
    return "\n".join(values)


def source_sections(row):
    if row["origin"].endswith(".docx"):
        return word_to_sections(ROOT / row["origin"], row["owner_department"]), ""
    number = int(Path(row["file"]).name.split("_", 1)[0])
    spec = next(item for item in NEW_DOCUMENTS if item["number"] == number)
    return (spec["title"], spec["sections"]), f"文档编号：{spec['code']}；版本1.0；2026年10月1日生效。"


def author_word(target, row):
    (title, sections), code = source_sections(row)
    blocks = [{"kind": "p", "text": title, "level": 0},
              {"kind": "p", "text": COMPANY, "level": 0},
              {"kind": "p", "text": metadata(row["group"], row["owner_department"]), "level": 0}]
    if code:
        blocks.append({"kind": "p", "text": code, "level": 0})
    blocks.append({"kind": "p", "text": NOTICE, "level": 0})
    for heading, content in sections:
        if heading:
            match = re.match(r"^(\d+(?:\.\d+)*)\s+(.+)$", heading)
            level = 2 if match and "." in match[1] else 1
            blocks.append({"kind": "p", "text": match[2] if match else heading, "level": level})
        if isinstance(content, list) and content and isinstance(content[0], list):
            blocks.append({"kind": "table", "rows": content})
        else:
            blocks.extend({"kind": "p", "text": str(value), "level": 0}
                          for value in (content if isinstance(content, list) else [content]))
    make_word(blocks, target)
    doc = Document(target)
    doc.core_properties.last_modified_by = COMPANY
    doc.core_properties.comments = ""
    for name in ("Title", "Heading 1", "Heading 2"):
        style = doc.styles[name]
        style.font.underline = False
        style.paragraph_format.keep_with_next = True
        for border in list(style.element.xpath("./w:pPr/w:pBdr")):
            border.getparent().remove(border)
    for table in doc.tables:
        is_metadata = table.cell(0, 0).text == "文档编号"
        for row_index, table_row in enumerate(table.rows):
            for column_index, cell in enumerate(table_row.cells):
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                header = row_index == 0 and not is_metadata
                label = is_metadata and column_index % 2 == 0
                if header or label:
                    shading = OxmlElement("w:shd")
                    shading.set(qn("w:fill"), "EEF2F6" if header else "F3F6F8")
                    cell._tc.get_or_add_tcPr().append(shading)
                for paragraph in cell.paragraphs:
                    paragraph.paragraph_format.widow_control = True
                    # Rows may paginate, but don't glue a whole table together.
                    paragraph.paragraph_format.keep_with_next = False
                    short_value = len(cell.text) <= 12 and re.fullmatch(r"[\d.%年月日元至]+", cell.text)
                    paragraph.alignment = (WD_ALIGN_PARAGRAPH.CENTER if header or short_value
                                           else WD_ALIGN_PARAGRAPH.LEFT)
                    for run in paragraph.runs:
                        run.font.color.rgb = RGBColor(0, 0, 0)
                        run.bold = header or label
    doc.save(target)
    return blocks


def content_check(target, original_pdf, blocks):
    """Compare every business paragraph/cell with the existing PDF and DOCX."""
    pdf_text = "\n".join(page.extract_text() or "" for page in PdfReader(original_pdf).pages)
    pdf_text = re.sub(r"第\s*\d+\s*页", "", pdf_text)
    pdf_text = pdf_text.replace(f"{COMPANY} | 内部资料", "").replace("虚构测试资料", "")
    actual_text = word_text(target)
    anchors = []
    for block in blocks:
        if block["kind"] == "table":
            anchors.extend(value for values in block["rows"] for value in values)
        elif block["text"] not in (COMPANY, NOTICE) and not block["text"].startswith("文档编号："):
            anchors.append(block["text"])
    missing_from_pdf = [text for text in anchors if normalized(text) not in normalized(pdf_text)]
    missing_from_word = [text for text in anchors if normalized(text) not in normalized(actual_text)]
    assert not missing_from_pdf, (original_pdf, missing_from_pdf)
    assert not missing_from_word, (target, missing_from_word)
    doc = Document(target)
    assert doc.paragraphs[0].style.name == "Title"
    assert any(paragraph.style.name == "Heading 1" for paragraph in doc.paragraphs)
    assert doc.tables
    assert all(table.autofit is False for table in doc.tables)
    assert all(cell.vertical_alignment == WD_CELL_VERTICAL_ALIGNMENT.CENTER
               for table in doc.tables for row in table.rows for cell in row.cells)
    assert all(not table._tbl.xpath(".//w:trHeight[@w:hRule='exact']") for table in doc.tables)
    assert all(table._tbl.xpath("./w:tblPr/w:tblBorders") for table in doc.tables)
    assert not doc.styles["Title"].element.xpath("./w:pPr/w:pBdr")
    assert all(not paragraph._p.xpath("./w:pPr/w:pBdr") for paragraph in doc.paragraphs)
    with zipfile.ZipFile(target) as bundle:
        assert bundle.testzip() is None
    return dict(document=target.name, business_anchors=len(anchors), content_preserved=True,
                native_tables=len(doc.tables), native_headings=True, structure="passed",
                rendered_pages=None, layout_review="pending_user_manual_review")


def build(output):
    if output.exists() and any(output.iterdir()):
        raise FileExistsError(f"Choose a new empty output folder: {output}")
    manifest = json.loads((SOURCE / "manifest.json").read_text(encoding="utf-8"))
    assert manifest["edition"] == "department-offline-v2" and manifest["total_documents"] == 30
    original_hashes = {row["file"]: row["sha256"] for row in manifest["documents"]}
    for relative, expected in original_hashes.items():
        assert sha256(SOURCE / relative) == expected, relative
    # Copy only the 24 retained files; replacement PDFs stay in the old edition.
    output.mkdir(parents=True, exist_ok=True)
    mapping, word_checks = {}, []
    for row in manifest["documents"]:
        source = SOURCE / row["file"]
        number = int(source.name.split("_", 1)[0])
        target = output / row["file"]
        target.parent.mkdir(parents=True, exist_ok=True)
        if number not in REPLACE_NUMBERS:
            shutil.copy2(source, target)
            continue
        assert source.suffix == ".pdf"
        target = target.with_suffix(".docx")
        blocks = author_word(target, row)
        word_checks.append(content_check(target, source, blocks))
        mapping[source.name] = target.name
        text = word_text(target)
        row.update(replaces_file=row["file"], replaces_sha256=row["sha256"],
                   file=target.relative_to(output).as_posix(), format="docx", sha256=sha256(target),
                   bytes=target.stat().st_size, text_characters=len(text), source_characters=len(text),
                   layout_review="pending_user_manual_review",
                   features=["Word原生标题", "Word原生表格", "部门目录", "人工排版检查待完成"])
        row.pop("pages", None)
        row.pop("embedded_text_characters", None)
    assert len(mapping) == 6
    manifest.update(edition="department-offline-v3-docx-manual-review", replaced_pdf_count=6,
                    based_on_edition="department-offline-v2", manual_docx_layout_review_accepted=True,
                    counts=dict(Counter(row["format"] for row in manifest["documents"])),
                    source_characters=sum(row["source_characters"] for row in manifest["documents"]))
    assert manifest["counts"] == EXPECTED_COUNTS
    assert Counter(row["group"] for row in manifest["documents"]) == manifest["group_counts"]
    assert len({row["sha256"] for row in manifest["documents"]}) == 30
    dump_json(output / "manifest.json", manifest)
    with (output / "入库清单.csv").open("w", encoding="utf-8-sig", newline="") as stream:
        writer = csv.writer(stream)
        keys = ("file", "group", "owner_department", "target_knowledge_base", "visibility", "status", "format", "sha256")
        writer.writerow(["相对路径", "资料分组", "归口部门", "建议目标知识库", "可见范围", "状态", "格式", "SHA256"])
        writer.writerows([row[key] for key in keys] for row in manifest["documents"])
    evaluation = output / "evaluation"
    evaluation.mkdir()
    for name in ("questions.jsonl", "conversation_cases.json", "permission_cases.json", "scan_ground_truth.json"):
        text = (SOURCE / "evaluation" / name).read_text(encoding="utf-8")
        for before, after in mapping.items():
            text = text.replace(before, after)
        (evaluation / name).write_text(text, encoding="utf-8")
    dump_json(evaluation / "docx_content_checks.json", dict(documents=word_checks, business_content="passed",
              visual_layout="not_checked_by_agent", manual_review_accepted_by_user=True))
    old_checks = json.loads((SOURCE / "evaluation/generation_checks.json").read_text(encoding="utf-8"))
    dump_json(evaluation / "generation_checks.json", dict(document_counts="passed", source_hashes="passed",
              unique_output_hashes="passed", department_counts="passed", docx_content_checks="passed",
              docx_visual_review="pending_user_manual_review; user explicitly accepted this delivery mode",
              word_rendering="not_run: missing LibreOffice soffice.exe; user will review in Word or WPS",
              pdf_visual_review=dict(status="inherited_from_v2_for_byte_identical_files", documents=12,
                                     pages=sum(row.get("pages", 0) for row in manifest["documents"])),
              actual_ocr=old_checks["actual_ocr"], parser_validation="pending",
              retrieval_and_model_quality="not_run", external_model_calls=0))
    readme = (SOURCE / "README.md").read_text(encoding="utf-8")
    readme = readme.replace("18份PDF、8份Markdown、4份TXT", "6份DOCX、12份PDF、8份Markdown、4份TXT")
    readme = readme.replace("新包替代原20份包用于测试", "v3包替代v2部门版或原20份包用于测试")
    readme = readme.replace("旧工具和旧20份包均保留。", "旧工具、v2部门版和原20份包均保留。")
    old_word_note = "按照用户选择，本次不新增DOCX。原包6份Word的内容重新排版为PDF，并检查PDF页图；这些PDF不是Word实际分页的证明。"
    readme = readme.replace(old_word_note, "本版将4份制度流程和2份操作规范的PDF替换为原生DOCX，保持业务内容与部门分类。用户已同意跳过自动DOCX分页验收，由用户在Word或WPS中逐页检查；DOCX内容和结构检查不能证明分页排版正确。")
    readme = re.sub(r"生成工具可运行.*", "生成工具：使用文档运行时执行 `scripts/replace_xinghai_corpus_docx.py --accept-manual-layout-review`，运行项目切分验证器后再执行该工具的 `--package-only`。生成器拒绝覆盖非空目录。v2包保留，本版未新增或修改PDF。", readme)
    for before, after in mapping.items():
        readme = readme.replace(before, after)
    readme = readme.replace("## 上传操作", "## DOCX人工验收\n\n打开本目录的 `人工排版检查清单.md`，使用Word或WPS打印布局逐页检查6份DOCX。不要把检查清单上传入库。文件保留原部门目录；不能将同内容的旧PDF和新DOCX同时上传同一知识库。\n\n## 上传操作")
    (output / "README.md").write_text(readme, encoding="utf-8")
    checklist = "# DOCX人工排版检查清单\n\n本版包含6份待人工排版验收的DOCX。请用Word或WPS打开，切换打印布局或打印预览，逐页查看；目前尚未确认实际页数。\n\n"
    for row in manifest["documents"]:
        if row["format"] == "docx":
            checklist += f"- [ ] {row['group']}：{Path(row['file']).name}\n"
    checklist += "\n## 每份文档的检查项\n\n- 标题、归口部门、可见范围、编号和版本正确。\n- 文字完整，无重叠、乱码或超出页边距。\n- 表格列宽合理，单元格文字不截断，跨页表格表头重复。\n- 标题没有孤立在页底，没有多余空白页或明显不合理分页。\n- 排版修改后重新检查；不要仅凭文件能打开就勾选完成。\n\n## 检查后入库\n\n只上传documents中的业务文件，不上传本清单、README、入库清单或evaluation。若修改了DOCX，文件哈希会变化，应重新记录清单和执行切分检查；本次不自动上传或替换数据库中的旧文件。\n"
    (output / "人工排版检查清单.md").write_text(checklist, encoding="utf-8")
    assert all(sha256(SOURCE / relative) == expected for relative, expected in original_hashes.items())
    print(json.dumps(dict(total_documents=30, counts=manifest["counts"], original_v2_unchanged=True,
                         business_content="passed", layout_review="pending_user_manual_review")))


def package(output):
    manifest = json.loads((output / "manifest.json").read_text(encoding="utf-8"))
    parsed = json.loads((output / "evaluation/splitter_checks.json").read_text(encoding="utf-8"))
    assert manifest["counts"] == EXPECTED_COUNTS
    assert parsed["hashes"] == "passed" and parsed["parsed_documents"] == 29
    assert parsed["single_turn_cases"] == 45 and not parsed["exact_evidence_misses"]
    assert parsed["cross_page_condition_preserved"] and parsed["dual_column_condition_preserved"]
    assert all(sha256(output / row["file"]) == row["sha256"] for row in manifest["documents"])
    checks_path = output / "evaluation/generation_checks.json"
    checks = json.loads(checks_path.read_text(encoding="utf-8"))
    checks["parser_validation"] = {key: value for key, value in parsed.items()
                                  if key not in ("documents", "evidence_checks")}
    dump_json(checks_path, checks)
    archive = output.parent / f"{COMPANY}_部门知识库30份_v3.zip"
    with zipfile.ZipFile(archive, "w", compression=zipfile.ZIP_DEFLATED) as bundle:
        for path in sorted(output.rglob("*")):
            if path.is_file():
                bundle.write(path, Path(output.name) / path.relative_to(output))
    with zipfile.ZipFile(archive) as bundle:
        assert bundle.testzip() is None
        documents = [name for name in bundle.namelist() if "/documents/" in name]
        assert len(documents) == 30
        assert all(hashlib.sha256(bundle.read(output.name + "/" + row["file"])).hexdigest() == row["sha256"]
                   for row in manifest["documents"])
    print(json.dumps(dict(archive=str(archive), zip_integrity="passed", documents=30,
                         counts=manifest["counts"], parsed_documents=parsed["parsed_documents"],
                         chunks=parsed["chunks"], docx_layout_review="pending_user_manual_review")))


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--output", type=Path, default=OUTPUT)
    parser.add_argument("--accept-manual-layout-review", action="store_true")
    parser.add_argument("--package-only", action="store_true")
    args = parser.parse_args()
    if args.package_only:
        package(args.output.resolve())
    elif args.accept_manual_layout_review:
        build(args.output.resolve())
    else:
        parser.error("DOCX authoring requires explicit acceptance of manual layout review.")
