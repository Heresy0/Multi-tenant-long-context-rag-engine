"""Build a small, offline fictional Chinese enterprise corpus (no model calls).

Run with the bundled document Python runtime. This does not upload or index files.
Original seed documents remain unchanged. Review files are kept outside documents/.
"""
from __future__ import annotations

import argparse
import hashlib
import json
import re
import subprocess
import zipfile
from collections import Counter
from pathlib import Path
from xml.sax.saxutils import escape

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from docx.table import Table as WordTable
from docx.text.paragraph import Paragraph as WordParagraph
from PIL import Image, ImageDraw, ImageFont
from pypdf import PdfReader
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.pdfgen import canvas
from reportlab.platypus import (
    CondPageBreak, LongTable, Paragraph, SimpleDocTemplate, TableStyle,
)

ROOT = Path(__file__).resolve().parents[1]
COMPANY = "星海科技有限公司"
NOTICE = "虚构测试资料，仅用于知识库开发与评估，不具有实际管理或法律效力。"
FONT_PATH = Path("C:/Windows/Fonts/msyh.ttc")
FONT_NAME = "XinghaiCJK"
PAGE_W, PAGE_H = A4
MARGIN = 48
BODY_WIDTH = PAGE_W - 2 * MARGIN
SEED_DIR = ROOT / "sample_docs/fictional_enterprise"


def rebrand(text: str) -> str:
    text = text.replace("启明数字科技有限公司", COMPANY)
    text = text.replace("启明开放平台", "星海开放平台")
    text = text.replace("QM-", "XH-").replace("X-QM-", "X-XH-")
    # Align inherited FAQ/table wording with the authoritative HR policy.
    text = text.replace("新员工按在岗月数折算", "新员工按剩余日历天数折算")
    text = text.replace("当年度按在岗月数折算", "当年度按剩余日历天数折算")
    if text.startswith("A  金额超过50000元至300000元时"):
        text = (
            "A  恰好50000元属于5000元以上至50000元档，至少取得2家有效报价，"
            "由部门负责人批准并形成比价记录。超过50000元至300000元时，"
            "至少取得3家有效报价并完成供应商准入。单一来源另行共同审批。"
        )
    return text


def iter_blocks(doc):
    for element in doc.element.body:
        if element.tag == qn("w:p"):
            yield WordParagraph(element, doc)
        elif element.tag == qn("w:tbl"):
            yield WordTable(element, doc)


def read_seed(index: int) -> tuple[Path, list[dict]]:
    path = next(SEED_DIR.glob(f"{index:02d}_*.docx"))
    blocks = []
    for block in iter_blocks(Document(path)):
        if isinstance(block, WordTable):
            blocks.append({"kind": "table", "rows": [
                [rebrand(c.text) for c in row.cells] for row in block.rows
            ]})
        elif block.text.strip():
            level = 0
            match = re.match(r"^(\d+(?:\.\d+)*)\s+", block.text)
            if match and len(block.text) < 65:
                level = match.group(1).count(".") + 1
            blocks.append({"kind": "p", "text": rebrand(block.text), "level": level})
    return path, blocks


def set_font(style, size: float):
    style.font.name = "Microsoft YaHei"
    style.font.size = Pt(size)
    style.font.color.rgb = RGBColor(0, 0, 0)
    style.element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "微软雅黑")


def make_word(blocks: list[dict], target: Path):
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Inches(8.5), Inches(11)
    section.top_margin = section.bottom_margin = Inches(0.8)
    section.left_margin = section.right_margin = Inches(0.9)
    for name, size in (("Normal", 11), ("Title", 22), ("Heading 1", 15), ("Heading 2", 12)):
        style = doc.styles[name]
        set_font(style, size)
        style.paragraph_format.space_after = Pt(7 if name == "Normal" else 10)
        style.paragraph_format.line_spacing = 1.2
        if name.startswith("Heading"):
            style.paragraph_format.space_before = Pt(12)
            style.paragraph_format.keep_with_next = True
    doc.core_properties.title = blocks[0]["text"]
    doc.core_properties.author = COMPANY
    doc.core_properties.subject = NOTICE
    for i, block in enumerate(blocks):
        if block["kind"] == "p":
            style = "Title" if i == 0 else f"Heading {min(block['level'], 2)}" if block["level"] else "Normal"
            p = doc.add_paragraph(block["text"], style)
            p.paragraph_format.widow_control = True
            continue
        rows = block["rows"]
        count = len(rows[0])
        table = doc.add_table(rows=len(rows), cols=count)
        table.autofit = False
        metadata = rows[0][0] == "文档编号"
        widths = [1.0, 2.35, 1.0, 2.35] if metadata else {
            2: [1.6, 5.1], 3: [1.4, 2.6, 2.7], 4: [1.1, 2.0, 1.6, 2.0],
        }.get(count, [6.7 / count] * count)
        for column, width in zip(table.columns, widths):
            column.width = Inches(width)
        properties = table._tbl.tblPr
        borders = OxmlElement("w:tblBorders")
        for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
            border = OxmlElement(f"w:{edge}")
            for key, value in (("val", "single"), ("sz", "4"), ("color", "D9D9D9")):
                border.set(qn(f"w:{key}"), value)
            borders.append(border)
        properties.append(borders)
        for row_index, (row, values) in enumerate(zip(table.rows, rows)):
            row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
            if row_index == 0 and not metadata:
                row._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
            for column_index, (cell, value) in enumerate(zip(row.cells, values)):
                cell.width = Inches(widths[column_index])
                cell.text = value
                tcpr = cell._tc.get_or_add_tcPr()
                padding = OxmlElement("w:tcMar")
                for edge, amount in (("top", 90), ("bottom", 90), ("left", 100), ("right", 100)):
                    item = OxmlElement(f"w:{edge}")
                    item.set(qn("w:w"), str(amount))
                    item.set(qn("w:type"), "dxa")
                    padding.append(item)
                tcpr.append(padding)
                for p in cell.paragraphs:
                    p.paragraph_format.space_after = Pt(2)
                    p.paragraph_format.line_spacing = 1.15
                    for run in p.runs:
                        run.font.size = Pt(10)
                        run.bold = row_index == 0 and not metadata
        doc.add_paragraph().paragraph_format.space_after = Pt(1)
    doc.save(target)


def seed_to_markdown(blocks: list[dict]) -> str:
    lines = []
    for i, block in enumerate(blocks):
        if block["kind"] == "table":
            rows = [[v.replace("|", "\\|").replace("\n", "<br>") for v in row] for row in block["rows"]]
            lines.append("| " + " | ".join(rows[0]) + " |")
            lines.append("| " + " | ".join(["---"] * len(rows[0])) + " |")
            lines.extend("| " + " | ".join(row) + " |" for row in rows[1:])
        else:
            text = block["text"]
            level = 1 if i == 0 else min(block["level"] + 1, 4) if block["level"] else 0
            if re.match(r"Q\d+\s", text):
                level = 3
            lines.append("#" * level + " " + text if level else text)
        lines.append("")
    return "\n".join(lines).strip() + "\n"


def paragraph(text: str, kind: str = "body") -> Paragraph:
    styles = {
        "body": ParagraphStyle("body", fontName=FONT_NAME, fontSize=11, leading=18,
                               wordWrap="CJK", spaceAfter=9),
        "title": ParagraphStyle("title", fontName=FONT_NAME, fontSize=22, leading=30, spaceAfter=16),
        "heading": ParagraphStyle("heading", fontName=FONT_NAME, fontSize=14, leading=22,
                                  spaceBefore=14, spaceAfter=9, keepWithNext=True),
        "small": ParagraphStyle("small", fontName=FONT_NAME, fontSize=9, leading=14,
                                wordWrap="CJK", spaceAfter=8),
        "cell": ParagraphStyle("cell", fontName=FONT_NAME, fontSize=9.5, leading=15, wordWrap="CJK"),
    }
    return Paragraph(escape(text).replace("\n", "<br/>"), styles[kind])


def page_furniture(c, doc):
    c.saveState()
    c.setFont(FONT_NAME, 8)
    c.setFillColor(colors.HexColor("#666666"))
    c.drawString(MARGIN, PAGE_H - 28, f"{COMPANY} | 内部资料")
    c.drawString(MARGIN, 25, "虚构测试资料")
    c.drawRightString(PAGE_W - MARGIN, 25, f"第 {doc.page} 页")
    c.restoreState()


def pdf_table(rows: list[list[str]], widths=None):
    if widths is None:
        widths = [BODY_WIDTH / len(rows[0])] * len(rows[0])
    table = LongTable([[paragraph(v, "cell") for v in row] for row in rows],
                      colWidths=widths, repeatRows=1, hAlign="LEFT")
    table.setStyle(TableStyle([
        ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#D9D9D9")),
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#EEF2F6")),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("LEFTPADDING", (0, 0), (-1, -1), 9), ("RIGHTPADDING", (0, 0), (-1, -1), 9),
        ("TOPPADDING", (0, 0), (-1, -1), 8), ("BOTTOMPADDING", (0, 0), (-1, -1), 8),
    ]))
    return table


def make_pdf(target: Path, title: str, code: str, sections: list[tuple[str, object]], status="现行"):
    story = [paragraph(title, "title"), paragraph(COMPANY),
             paragraph(f"文档编号：{code} | 状态：{status}", "small"), paragraph(NOTICE, "small")]
    for heading, content in sections:
        if isinstance(content, list) and content and isinstance(content[0], list):
            # Keep room for the title/header/first row, not the entire long table.
            story.append(CondPageBreak(140))
            heading_paragraph = paragraph(heading, "heading")
            heading_paragraph.keepWithNext = False
            story.append(heading_paragraph)
            widths = [BODY_WIDTH * fraction for fraction in (0.24, 0.18, 0.14, 0.44)] if len(content) > 15 else None
            story.append(pdf_table(content, widths))
        else:
            story.append(paragraph(heading, "heading"))
            for text in content if isinstance(content, list) else [content]:
                story.append(paragraph(str(text)))
    SimpleDocTemplate(str(target), pagesize=A4, leftMargin=MARGIN, rightMargin=MARGIN,
                      topMargin=52, bottomMargin=48, title=title, author=COMPANY,
                      subject=NOTICE).build(story, onFirstPage=page_furniture, onLaterPages=page_furniture)


def draw_wrapped(c, text: str, x: float, top: float, width: float, size=11, leading=19):
    p = Paragraph(escape(text).replace("\n", "<br/>"),
                  ParagraphStyle("canvas-body", fontName=FONT_NAME, fontSize=size,
                                 leading=leading, wordWrap="CJK"))
    _, height = p.wrap(width, PAGE_H)
    if top - height < 50:
        raise ValueError(f"Content would extend outside page: {text[:30]}")
    p.drawOn(c, x, top - height)
    return top - height


def canvas_page(c, title, code, page):
    c.setFont(FONT_NAME, 8)
    c.setFillColor(colors.HexColor("#666666"))
    c.drawString(MARGIN, PAGE_H - 28, f"{COMPANY} | 内部资料")
    c.drawString(MARGIN, 25, "虚构测试资料")
    c.drawRightString(PAGE_W - MARGIN, 25, f"第 {page} 页")
    c.setFillColor(colors.black)
    top = draw_wrapped(c, title, MARGIN, PAGE_H - 65, BODY_WIDTH, 19, 28)
    return draw_wrapped(c, code, MARGIN, top - 14, BODY_WIDTH, 9, 15) - 25


def cross_page_pdf(target: Path) -> str:
    c = canvas.Canvas(str(target), pagesize=A4)
    c.setTitle("客户数据外发审批补充规定")
    c.setAuthor(COMPANY)
    c.setSubject(NOTICE)
    top = canvas_page(c, "客户数据外发审批补充规定", "XH-SEC-2026-09 | 版本1.0 | 2026年9月1日生效", 1)
    text1 = (
        "1 适用范围\n本规定是XH-SEC-2026-02的补充。适用于交付、客服和研发团队向已批准的"
        "客户协作空间交付故障分析材料，不改变机密数据的基本审批要求。\n\n"
        "2 提交前检查\n申请人须核实接收方身份、合同依据、数据范围、保存期限和删除方式。"
        "材料应采用最小化原则，去除令牌、生产密钥、真实身份证号和银行卡号。\n\n"
        "3 审批记录\n系统应保留数据所有者、信息安全负责人、申请编号和批准时间。"
        "口头同意、截图转述及历史工单不能替代当次书面批准。"
    )
    draw_wrapped(c, text1, MARGIN, top, BODY_WIDTH)
    lead = "客户故障排查确需外发脱敏日志时，申请人须先取得数据所有者与信息安全负责人的书面批准，并且"
    draw_wrapped(c, lead, MARGIN, 116, BODY_WIDTH)
    c.showPage()
    # No repeated title between the two halves: the logical paragraph crosses pages.
    c.setFont(FONT_NAME, 8)
    c.setFillColor(colors.HexColor("#666666"))
    c.drawString(MARGIN, PAGE_H - 28, f"{COMPANY} | 内部资料")
    c.drawString(MARGIN, 25, "虚构测试资料")
    c.drawRightString(PAGE_W - MARGIN, 25, "第 2 页")
    c.setFillColor(colors.black)
    tail = "仅可通过批准的加密协作空间发送，接收方下载权限必须在72小时后自动失效；未经批准不得改用普通邮件附件。"
    top = draw_wrapped(c, tail, MARGIN, PAGE_H - 68, BODY_WIDTH)
    draw_wrapped(c,
        "4 例外与关闭\n如客户无法使用批准空间，由数据所有者与信息安全负责人共同批准替代通道，"
        "并明确有效期及审计方式。申请人不得自行选择公共网盘。\n\n"
        "完成交付后，申请人应验证权限关闭并在申请单中附上关闭时间。文件接收并不等于允许长期保存。\n\n"
        "5 责任归属\n数据所有者确认业务必要性；信息安全负责人确认通道和脱敏控制；申请人负责执行和留存证据。",
        MARGIN, top - 28, BODY_WIDTH)
    c.save()
    return lead + tail


def two_column_pdf(target: Path):
    c = canvas.Canvas(str(target), pagesize=A4)
    c.setTitle("生产运维值班与回退手册")
    c.setAuthor(COMPANY)
    c.setSubject(NOTICE)
    top = canvas_page(c, "生产运维值班与回退手册", "XH-OPS-2026-10 | 版本1.0 | 2026年9月1日生效", 1)
    top = draw_wrapped(c, "本手册适用于星海协作云的生产值班。左栏说明日常监控，右栏说明事故处置，均由当班人员执行。",
                       MARGIN, top, BODY_WIDTH) - 28
    gap = 28
    width = (BODY_WIDTH - gap) / 2
    left = (
        "1 日常值班\n每个班次交接时确认告警订阅、发布日历、未关闭工单及主备值班人。"
        "主值班负责告警确认，备值班负责验证和记录。\n\n"
        "1.1 监控项目\n每5分钟检查核心API成功率、P95延迟、消息队列积压和数据库连接数。"
        "值班看板以租户维度展示，不在截图中包含客户业务正文。\n\n"
        "1.2 发布前检查\n发布负责人确认备份点、回退脚本、维护通知及验证清单。"
        "临时跳过检查必须在变更单中记录理由和批准人。\n\n"
        "1.3 归档\n值班交接记录保留180日，交接不改变未关闭工单的责任人。"
    )
    right = (
        "2 事故处置\nP1故障按技术手册要求15分钟内首次响应，每30分钟更新状态；响应不等于恢复。\n\n"
        "2.1 回退条件\n核心API成功率连续10分钟低于99.0%，或发现数据写入错误时，"
        "发布负责人立即决定是否回退。先停止新增流量，再切换上一个稳定版本。\n\n"
        "2.2 恢复验证\n恢复后验证登录、审批、查询和Webhook四条关键链路。"
        "不得仅因CPU恢复正常就关闭事件。\n\n"
        "2.3 复盘\nP1恢复后5个工作日内提供初步复盘摘要。未核实的根因写为待核实，不提前作结论。"
    )
    draw_wrapped(c, left, MARGIN, top, width)
    draw_wrapped(c, right, MARGIN + width + gap, top, width)
    c.save()


def scanned_pdf(target: Path, scratch: Path) -> str:
    lines = [COMPANY, "生产维护通知", "虚构测试资料", "通知编号：XH-MNT-2026-017",
             "发布日期：2026年9月18日", "维护窗口：2026年9月20日02:00至03:00（北京时间）",
             "影响范围：星海协作云生产环境的报表导出功能。", "登录、审批与开放API不受本次维护影响。",
             "维护负责人：林澈；验证负责人：周宁。", "维护期间导出任务暂停接收，已有任务保留。",
             "03:00后恢复接收，无须重新提交已排队任务。", "如03:00仍未恢复，由值班组发布延期通知。",
             "取消本次维护须在窗口开始前由发布负责人批准。", "联系渠道：企业服务台维护公告栏目。",
             "XH-MNT-2026-017 / REPORT EXPORT ONLY"]
    image = Image.new("RGB", (1654, 2339), "#FAFAF7")
    draw = ImageDraw.Draw(image)
    for index, line in enumerate(lines):
        size = 62 if index == 1 else 42 if index == 0 else 34
        draw.text((125, 150 + index * 110), line, fill="#242424", font=ImageFont.truetype(str(FONT_PATH), size))
    # Mild scan artifacts; keep source legible rather than manufacturing unreadable noise.
    image = image.rotate(0.25, resample=Image.Resampling.BICUBIC, expand=False, fillcolor="#FAFAF7")
    image_path = scratch / "maintenance-source.png"
    image.save(image_path)
    c = canvas.Canvas(str(target), pagesize=A4)
    c.setTitle("生产维护通知扫描件")
    c.setAuthor(COMPANY)
    c.setSubject(NOTICE)
    c.drawImage(str(image_path), 0, 0, PAGE_W, PAGE_H)
    c.save()
    return "\n".join(lines)


NEW_MD = {
    "17_协作云部署与备份技术手册.md": """# 协作云部署与备份技术手册

星海科技有限公司

文档编号：XH-TECH-2026-11；版本：1.0；生效日期：2026年9月1日；状态：现行。

虚构测试资料，仅用于知识库开发与评估，不具有实际管理或法律效力。

## 1 部署边界

生产服务包含应用服务、PostgreSQL、对象存储和消息队列。沙箱与生产使用不同账户和数据库；沙箱不得接收真实的严格机密数据。域名均采用 example.invalid，仅用于示例。

### 1.1 配置检查

生产数据库必须使用专属账户与最小权限；访问凭据来自批准的密钥系统，不直接写入配置文件。以下片段只有非敏感示例，不是可运行的生产配置。

```yaml
service:
  name: xinghai-collaboration-cloud
  environment: sandbox
backup:
  schedule: "02:30 Asia/Shanghai"
  retention_days: 35
  restore_drill: monthly
```

## 2 备份与恢复目标

| 对象 | 备份方式 | 保留期 | 恢复目标 |
| --- | --- | --- | --- |
| PostgreSQL | 每日全量与连续日志归档 | 35日 | RPO 15分钟，RTO 2小时 |
| 上传附件 | 对象存储版本与每日增量清单 | 35日 | RPO 24小时，RTO 4小时 |
| 系统配置 | 每次已批准发布后归档 | 180日 | 恢复上一个稳定版本 |

RPO表示最多允许丢失的数据时间窗口，RTO表示目标恢复时长。恢复目标是本手册的内部运维目标，不替代客户订单中的服务承诺。

## 3 恢复演练

每月执行一次隔离环境恢复演练。平台研发部负责人确认恢复数据时间点、关键链路及租户隔离检查；恢复演练不得将客户数据导入公共测试服务。演练记录由值班组保留180日。

## 4 恢复操作注意事项

1. 冻结新增写入，记录最后成功事务时间。
2. 在隔离环境恢复，核对备份校验值和日志连续性。
3. 验证登录、审批、查询及Webhook后，再由事件负责人批准切流。
4. 保留事件记录，不因恢复完成而删除原始审计日志。
""",
    "18_海岚项目周报与变更记录.md": """# 海岚项目周报与变更记录

星海科技有限公司

文档编号：XH-PMO-2026-12；项目编号：XH-PRJ-026；报告日期：2026年9月18日；状态：已确认。

虚构测试资料，仅用于知识库开发与评估，不具有实际管理或法律效力。海岚制造有限公司及人员均为虚构。

## 1 本周进展

海岚项目为海岚制造有限公司交付审批与报表模块。项目经理林澈、客户项目负责人赵苒于2026年9月18日确认当前需求基线。项目不包含ERP数据迁移，不应从标题推断额外交付范围。

## 2 变更CR-026-03

客户要求新增跨部门审批代理功能，评估工作量为4人日，原定2026年9月25日的用户验收调整至2026年9月29日。该变更超过2人日且影响里程碑，已于2026年9月17日取得双方授权代表书面批准。

| 项目 | 变更前 | 变更后 |
| --- | --- | --- |
| 用户验收日期 | 2026年9月25日 | 2026年9月29日 |
| 代理审批功能 | 不在需求基线 | 纳入CR-026-03 |
| 额外工作量 | 0人日 | 4人日 |

## 3 风险与待办

代理审批的撤销权限测试由测试负责人沈禾在2026年9月23日前完成。客户验收账号由赵苒在2026年9月24日前确认。风险登记R-026-02为客户账号确认延迟，当前状态为开放；此周报未宣布风险已关闭。

## 4 本周决议

未批准的需求不得进入生产。最终验收情况以XH-PMO-2026-14验收纪要为准；本周报只记录截至报告日的计划，不证明未来验收已完成。
""",
}

NEW_TXT = {
    "19_生产支持工单与故障记录.txt": """生产支持工单与故障记录
星海科技有限公司
文档编号：XH-CS-2026-15
状态：已归档；记录截止：2026年9月22日
虚构测试资料，仅用于知识库开发与评估，不具有实际管理或法律效力。

1 工单INC-2026-0912
发现时间：2026年9月12日10:05（北京时间）。
等级：P2。影响：海岚制造沙箱应用的报表查询返回RATE-001及HTTP 429。
10:22支持工程师顾言首次响应，确认应用每分钟请求超过120次。
处置：按Retry-After等待，使用带抖动的指数退避；不得创建多个应用绕过限制。
10:48调用方修正客户端并恢复查询。根因：客户端并发配置未按默认限流调整，不是生产数据库故障。
结论：首次响应耗时17分钟，恢复耗时43分钟，二者不是同一个指标。

2 工单INC-2026-0920
关联通知：XH-MNT-2026-017。
02:15客户咨询生产报表导出任务状态。维护期间暂停接收新任务，既有任务保留。
03:00导出功能恢复接收，已排队任务继续执行。
客服未要求客户重新提交排队任务。登录、审批、开放API均不受此次计划维护影响。

3 故障信息提交要求
客户提交租户编号、环境、时间范围、X-Request-Id、接口路径和脱敏摘要。
不得在工单里提供访问令牌、client_secret或生产密钥。
已关闭工单不能作为未来相同错误唯一根因的证据，需要重新核实。
""",
    "20_九月内部通知与值班安排.txt": """九月内部通知与值班安排
星海科技有限公司
文档编号：XH-ADM-2026-16；发布日期：2026年9月18日；状态：现行。
虚构测试资料，仅用于知识库开发与评估，不具有实际管理或法律效力。

1 知识库培训安排
培训主题：来源核对与敏感信息脱敏。
时间：2026年9月24日15:00至16:30（北京时间）。
方式：线上培训；主讲：信息安全负责人许澜。
参加人员：交付、客服和平台研发团队；请在9月23日18:00前在培训系统确认报名。

2 九月维护值班
关联维护通知XH-MNT-2026-017，窗口为9月20日02:00至03:00。
主值班：林澈；备值班与恢复验证：周宁。
维护只影响生产报表导出功能，不影响登录、审批和开放API。
值班交接应说明未关闭事项，不在群聊中发送客户原始数据。

3 文档使用提醒
报销时使用2026年7月1日生效的XH-FIN-2026-03版本1.4。
2025年旧版差旅标准已归档，只适用于其有效期间的历史事项，不作为当前报销标准。
未提供的薪资、客户合同金额或真实联系方式请咨询归口部门，不根据历史通知推测。
""",
}


def dump_json(path: Path, value):
    path.write_text(json.dumps(value, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")


def build(output: Path, render: bool):
    if not FONT_PATH.exists():
        raise FileNotFoundError(f"Required local CJK font not found: {FONT_PATH}")
    pdfmetrics.registerFont(TTFont(FONT_NAME, str(FONT_PATH), subfontIndex=0))
    document_dir = output / "documents"
    qa_dir = output / "evaluation"
    scratch = ROOT / "tmp/xinghai-corpus-qa"
    for path in (document_dir, qa_dir, scratch):
        path.mkdir(parents=True, exist_ok=True)
    # Refuse to overwrite unrelated pre-existing outputs.
    build_tag = output / "manifest.json"
    if any(document_dir.iterdir()) and not build_tag.exists():
        raise FileExistsError(f"Unrecognized output folder: {output}")
    records = []
    extracted = {}

    def register(name, category, features, source=None, status="current", text=None):
        path = document_dir / name
        if text is None:
            if path.suffix == ".pdf":
                reader = PdfReader(path)
                text = "\n".join(p.extract_text() or "" for p in reader.pages)
            else:
                text = path.read_text(encoding="utf-8")
        extracted[name] = text
        record = {"file": "documents/" + name, "format": path.suffix[1:],
                  "category": category, "status": status, "features": features,
                  "bytes": path.stat().st_size, "sha256": hashlib.sha256(path.read_bytes()).hexdigest(),
                  "text_characters": len(text), "origin": source or "offline_authored_template"}
        if path.suffix == ".pdf":
            record["pages"] = len(PdfReader(path).pages)
            record["embedded_text_characters"] = sum(len(p.extract_text() or "") for p in PdfReader(path).pages)
        records.append(record)

    for index in (1, 2, 3, 4, 5, 7):
        source, blocks = read_seed(index)
        output_index = 6 if index == 7 else index
        name = f"{output_index:02d}_" + source.name.split("_", 1)[1]
        make_word(blocks, document_dir / name)
        content = "\n".join(b["text"] if b["kind"] == "p" else "\n".join("|".join(r) for r in b["rows"]) for b in blocks)
        register(name, "合同" if index == 7 else "制度与流程", ["标题层级", "原生表格"],
                 source.relative_to(ROOT).as_posix(), text=content)

    make_pdf(document_dir / "07_公司组织与产品服务说明.pdf", "公司组织与产品服务说明", "XH-ORG-2026-01", [
        ("1 企业与服务范围", ["星海科技有限公司是一家虚构的企业软件服务公司，提供星海协作云、开放平台及实施服务。组织规模为180人。本说明只用于测试，不代表任何真实企业。",
                            "服务对象是需要审批、报表与开放接口的中小型企业。星海协作云不提供薪资计算服务，客户如需此功能须另行评估。"]),
        ("2 组织与责任", [["部门", "人数", "主要责任"], ["平台研发部", "64", "产品研发、平台可靠性与API版本维护"],
                          ["项目交付部", "42", "需求基线、变更管理与用户验收"], ["客户服务部", "28", "支持工单与维护通知"],
                          ["销售与市场部", "22", "销售机会、订单与客户沟通"], ["综合管理部", "24", "人事、财务、采购及合规支持"]]),
        ("3 产品边界", "星海协作云包含审批、报表和权限管理模块；开放平台提供认证、查询、写入及Webhook接口。订单中的专属支持目标和额度优先于默认手册。"),
        ("4 文档责任", "人事与财务制度由相应归口部门维护；技术手册由平台研发部维护。文档编号、版本和生效日期同时记录，归档版本不得直接作为当前操作标准。"),
    ])
    register("07_公司组织与产品服务说明.pdf", "公司与产品", ["正文", "组织表格"])

    make_pdf(document_dir / "08_差旅住宿标准_2025归档版.pdf", "差旅住宿标准历史版本", "XH-FIN-2025-03 | 版本1.2", [
        ("1 版本有效期", "本文件有效期为2025年1月1日至2026年6月30日。2026年7月1日起被XH-FIN-2026-03版本1.4替代，仅保留历史查询用途。"),
        ("2 历史住宿限额", [["城市类别", "普通员工每晚", "部门负责人及以上每晚"], ["A类城市", "550元", "750元"],
                          ["B类城市", "450元", "600元"], ["其他城市", "350元", "450元"]]),
        ("3 适用提醒", "A类城市为北京、上海、深圳、广州，金额为含税费后的单间价格。当前普通员工A类城市住宿上限不是550元，请查询2026年现行制度；历史费用按当时有效版本核实。"),
    ], status="已归档，不适用于当前报销")
    register("08_差旅住宿标准_2025归档版.pdf", "历史财务制度", ["版本冲突", "归档状态", "表格"], status="archived")

    cross_page_pdf(document_dir / "09_客户数据外发审批补充规定.pdf")
    register("09_客户数据外发审批补充规定.pdf", "信息安全", ["重复页眉页脚", "跨页条件与例外"])

    systems = ["身份认证", "审批引擎", "报表导出", "开放API", "Webhook投递", "租户管理", "审计日志", "对象存储",
               "生产数据库", "消息队列", "监控告警", "密钥托管", "备份归档", "邮件通知", "企业服务台", "文档预览",
               "全文检索", "向量检索", "任务调度", "数据脱敏", "客户端SDK", "权限复核", "发布管理", "恢复演练",
               "工单路由", "证书管理", "流量网关", "接口配额", "培训平台", "知识运营"]
    ledger = [["编号/系统", "归口团队", "复核周期", "复核证据"]]
    for i, system in enumerate(systems, 1):
        owner = "平台研发部" if i <= 14 else "客户服务部" if i >= 25 else "项目交付部"
        period = "每月" if i in (1, 9, 12, 22, 26) else "每季度"
        evidence = f"核对{system}的授权记录、操作日志与未关闭风险，登记责任人和整改期限。"
        ledger.append([f"ASSET-{i:03d}\n{system}", owner, period, evidence])
    make_pdf(document_dir / "10_系统资产复核台账.pdf", "系统资产复核台账", "XH-OPS-2026-13", [
        ("1 复核规则", "本台账由运营管理部汇总，覆盖30个系统或控制域。复核周期仅针对本台账检查活动，不替代信息安全制度要求的机密数据访问季度复核。记录缺失须在5个工作日内补齐，不得将未复核项目标记为通过。"),
        ("2 资产记录", ledger),
        ("3 闭环要求", "未通过的项目进入风险清单；责任人提交整改证据后，由复核人员确认关闭。台账应保留复核日期和批准人，不以生成报告代替执行复核。"),
    ])
    register("10_系统资产复核台账.pdf", "资产与运营", ["多页长表", "重复表头", "精确编号"])

    two_column_pdf(document_dir / "11_生产运维值班与回退手册.pdf")
    register("11_生产运维值班与回退手册.pdf", "运维", ["双栏阅读顺序", "独立章节"])
    scan_text = scanned_pdf(document_dir / "12_生产维护通知_扫描件.pdf", scratch)
    register("12_生产维护通知_扫描件.pdf", "维护公告", ["中文图片型PDF", "OCR必需", "精确编号"], text=scan_text)

    make_pdf(document_dir / "13_海岚项目最终验收纪要.pdf", "海岚项目最终验收纪要", "XH-PMO-2026-14", [
        ("1 会议与项目", "项目编号XH-PRJ-026；客户为虚构的海岚制造有限公司。会议日期2026年9月29日，项目经理林澈、客户项目负责人赵苒及测试负责人沈禾参加。"),
        ("2 验收范围", "本次验收覆盖审批与报表模块，以及CR-026-03批准的跨部门审批代理功能。不包含ERP数据迁移。原计划9月25日验收，变更批准后调整至9月29日。"),
        ("3 验收结果", [["验证项", "结果", "确认依据"], ["核心审批与报表", "通过", "UAT-026-09测试记录"],
                         ["代理审批撤销权限", "通过", "2026年9月23日回归记录"], ["客户验收账号", "完成", "2026年9月24日客户确认"],
                         ["遗留一般缺陷DEF-026-08", "书面接受", "导出按钮提示文字优化，计划2026年10月9日修复"]]),
        ("4 决议", "客户于2026年9月29日签署最终验收，质保期按流程手册从最终验收后计算90日。遗留问题由沈禾跟踪，不影响本次验收；不应将计划修复日期当作已经修复。风险R-026-02在客户验收账号确认后关闭。"),
    ])
    register("13_海岚项目最终验收纪要.pdf", "项目交付", ["跨文档关联", "状态变化", "表格"])

    make_pdf(document_dir / "14_协作云版本发布与兼容说明.pdf", "协作云版本发布与兼容说明", "XH-TECH-2026-17", [
        ("1 当前发布", "协作云2.6.0于2026年9月16日发布。开放API仍使用/v1主版本路径，本次新增可选响应字段approval_agent，不改变既有字段类型。客户端应忽略未知可选字段。"),
        ("2 版本记录", [["版本", "发布日期", "主要变化"], ["2.5.2", "2026年8月20日", "修复报表分页重复显示"],
                        ["2.6.0", "2026年9月16日", "新增审批代理能力与可选响应字段approval_agent"]]),
        ("3 兼容性说明", "协作云产品版本2.6.0与API主版本/v1是不同维度，不能从产品版本推断API已升级为/v2。本次发布不涉及旧主版本停服；后续计划停服原则上提前180日通知。"),
        ("4 已知限制", "审批代理仅对已授权角色生效，管理员不得使用共享账户配置代理。产品提供此能力不代表所有客户项目都已纳入合同范围，应核对需求基线与变更单。"),
    ])
    register("14_协作云版本发布与兼容说明.pdf", "产品与技术", ["版本区分", "表格", "精确字段"])

    for index, number in ((6, 15), (8, 16)):
        source, blocks = read_seed(index)
        name = f"{number:02d}_" + source.stem.split("_", 1)[1] + ".md"
        body = seed_to_markdown(blocks)
        (document_dir / name).write_text(body, encoding="utf-8")
        register(name, "技术手册" if index == 6 else "FAQ", ["Markdown标题", "Markdown表格", "问答" if index == 8 else "错误代码"],
                 source.relative_to(ROOT).as_posix())
    for name, body in NEW_MD.items():
        (document_dir / name).write_text(body, encoding="utf-8")
        register(name, "技术手册" if name.startswith("17") else "项目周报", ["Markdown标题", "代码围栏" if name.startswith("17") else "版本与跨文档关联"])
    for name, body in NEW_TXT.items():
        (document_dir / name).write_text(body, encoding="utf-8")
        register(name, "工单" if name.startswith("19") else "通知", ["纯文本标题", "编号与时间线"])

    # Ground truth is outside the upload folder, preventing answer leakage.
    dump_json(qa_dir / "scan_ground_truth.json", {
        "document": "12_生产维护通知_扫描件.pdf", "text": scan_text,
        "purpose": "人工源文本用于核对OCR；不得当作OCR识别结果或入库资料",
    })
    cases = make_cases(extracted)
    (qa_dir / "questions.jsonl").write_text("".join(json.dumps(c, ensure_ascii=False) + "\n" for c in cases), encoding="utf-8")
    dump_json(qa_dir / "conversation_cases.json", [
        {"id": "XH-MT-01", "turns": [
            {"question": "海岚项目的用户验收原定什么时候？", "expected_answer": "2026年9月25日。"},
            {"question": "后来为什么推迟了？", "expected_answer": "CR-026-03新增跨部门审批代理，工作量4人日，影响里程碑。"},
            {"question": "那最后验收通过了吗？", "expected_answer": "2026年9月29日签署最终验收，遗留一般缺陷DEF-026-08被书面接受。"}],
         "sources": ["18_海岚项目周报与变更记录.md", "13_海岚项目最终验收纪要.pdf"]},
        {"id": "XH-MT-02", "turns": [
            {"question": "现在普通员工去上海住宿每晚最多报多少？", "expected_answer": "650元，依据现行版本1.4。"},
            {"question": "去年呢？", "expected_answer": "2025年历史版本1.2为550元。"},
            {"question": "部门负责人现在呢？", "expected_answer": "850元。"}],
         "sources": ["02_差旅与费用报销管理制度.docx", "08_差旅住宿标准_2025归档版.pdf"]},
        {"id": "XH-MT-03", "turns": [
            {"question": "9月20日维护会影响什么？", "expected_answer": "生产报表导出，窗口02:00至03:00北京时间。"},
            {"question": "那登录还能用吗？", "expected_answer": "登录、审批和开放API不受此次维护影响。"},
            {"question": "之前排队的任务要重提吗？", "expected_answer": "不需要，已有任务保留并在恢复后继续执行。"}],
         "sources": ["12_生产维护通知_扫描件.pdf", "19_生产支持工单与故障记录.txt", "20_九月内部通知与值班安排.txt"]},
    ])

    expected_counts = {"docx": 6, "pdf": 8, "md": 4, "txt": 2}
    assert dict(Counter(r["format"] for r in records)) == expected_counts
    assert len(set(r["sha256"] for r in records)) == 20
    assert not PdfReader(document_dir / "12_生产维护通知_扫描件.pdf").pages[0].extract_text()
    assert len(PdfReader(document_dir / "09_客户数据外发审批补充规定.pdf").pages) == 2
    assert len(PdfReader(document_dir / "10_系统资产复核台账.pdf").pages) >= 2
    for name, text in extracted.items():
        assert COMPANY in text, name
        assert "启明" not in text and "QM-" not in text, name
    manifest = {"company": COMPANY, "fictional": True, "edition": "small-offline-v1",
                "counts": expected_counts, "total_documents": 20,
                "external_model_calls": 0, "source_characters": sum(len(t) for t in extracted.values()),
                "single_turn_questions": len(cases), "multi_turn_conversations": 3,
                "scale_note": "小规模功能测试集，不等同于中小企业真实文档规模或性能压测语料。",
                "documents": sorted(records, key=lambda r: r["file"])}
    dump_json(output / "manifest.json", manifest)
    dump_json(qa_dir / "generation_checks.json", {
        "document_counts": "passed", "unique_content_hashes": "passed", "rebranding": "passed",
        "single_turn_evidence_presence": "passed", "image_only_scan": "passed",
        "cross_page_pdf": "2 pages", "long_table_pdf": "multiple pages",
        "docx_visual_review": "pending: packaged renderer requires LibreOffice, not present in current bundled Windows runtime",
        "pdf_visual_review": "pending: inspect the rendered page PNGs before accepting",
        "ocr_runtime_test": "pending: image-only structure checked, actual OCR not yet verified",
        "rag_answer_quality": "not run; no upload, indexing, embeddings or model calls",
    })
    if render:
        for path in sorted(document_dir.glob("*.pdf")):
            pages_dir = scratch / path.stem
            pages_dir.mkdir(exist_ok=True)
            subprocess.run(["pdftoppm", "-r", "110", "-png", str(path), str(pages_dir / "page")], check=True)
    write_readme(output, manifest)
    package(output)
    print(json.dumps({"folder": str(output), "zip": str(output.parent / (COMPANY + "_知识库测试集.zip")),
                      "documents": 20, "counts": expected_counts,
                      "pdf_pages": sum(r.get("pages", 0) for r in records),
                      "source_characters": manifest["source_characters"], "questions": len(cases),
                      "rendered_pages": str(scratch) if render else "not rendered"}, ensure_ascii=False))


def make_cases(texts: dict[str, str]) -> list[dict]:
    cases = []

    def add(question, answer, facts, document, evidence, category="事实查询", difficulty="easy"):
        source = next(name for name in texts if name.startswith(document + "_"))
        # Table evidence uses the same pipe-separated form as the project splitter.
        normalize = lambda value: "".join(value.replace("|", "").split())
        assert normalize(evidence) in normalize(texts[source]), (source, evidence)
        cases.append({"id": f"XH-{len(cases) + 1:03d}", "question": question,
                      "reference_answer": answer, "answerable": True, "required_facts": facts,
                      "expected_evidence": [{"document": source, "text": evidence}],
                      "category": category, "difficulty": difficulty})

    add("星海员工每月可以有几次免说明的轻微迟到？", "每月2次，每次不超过15分钟。", ["2次", "15分钟"], "01", "每月允许2次、每次不超过15分钟")
    add("连续休4个工作日年休假要提前多久申请？", "提前5个工作日，连续超过3个工作日还需部门负责人审批。", ["5个工作日", "部门负责人"], "01", "超过3天提前5个工作日", "表格条件", "medium")
    add("现在普通员工去上海出差，住宿每晚报销上限多少？", "650元，现行版本1.4于2026年7月1日生效，旧版550元不可用于当前标准。", ["650元", "1.4"], "02", "每晚不超过650元", "版本冲突", "hard")
    add("部门负责人现在去A类城市的每晚住宿上限多少？", "每晚850元，含税费后的单间价格。", ["850元"], "02", "每晚不超过850元", "表格查询")
    add("行程结束后多久提交差旅报销？", "10个自然日内。", ["10个自然日"], "02", "行程结束或费用发生后10个自然日内提交报销单")
    add("发生客户名单误发，多久报告？", "立即停止扩大影响的操作，30分钟内报告，不自行删除日志。", ["30分钟", "不得自行删除日志"], "03", "并在30分钟内通过安全热线或事件系统报告", "安全流程")
    add("互联网暴露系统的严重漏洞原则上多久修复？", "72小时内。", ["72小时"], "03", "严重漏洞原则上72小时内修复", "表格查询")
    add("恰好50000元的采购至少需要几家报价，谁批准？", "至少2家有效报价，由部门负责人批准并形成比价记录。", ["2家", "部门负责人"], "04", "5000元以上至50000元|至少2家有效报价|部门负责人", "金额边界", "hard")
    add("6万元的一般采购需要几家报价和哪些批准人？", "至少3家有效报价，由部门负责人和财务负责人批准，并完成供应商准入。", ["3家", "部门负责人", "财务负责人"], "04", "50000元以上至300000元|至少3家有效报价|部门负责人和财务负责人", "金额边界", "medium")
    add("影响里程碑的变更谁批准？", "双方授权代表书面批准。", ["双方授权代表", "书面批准"], "05", "须由双方授权代表书面批准")
    add("默认项目质保期多长？", "最终验收后90日，除合同另有约定。", ["最终验收后90日"], "05", "项目质保期为最终验收后90日")
    add("月度可用性99.2%对应多少服务抵扣，何时申请？", "当月服务费10%，次月10个工作日内申请。", ["10%", "次月10个工作日"], "06", "低于99.5%但不低于99.0%|当月服务费的10%|次月10个工作日内", "区间计算", "medium")
    add("订阅终止后客户有多久导出数据？", "30日内可使用导出功能，窗口结束后启动删除。", ["30日"], "06", "订阅终止后30日内")
    add("星海公司组织规模多少人？", "180人。", ["180人"], "07", "组织规模为180人")
    add("平台研发部有多少人？", "64人。", ["64"], "07", "平台研发部|64|产品研发、平台可靠性与API版本维护", "表格查询")
    add("2025年普通员工在上海住宿每晚限额多少？", "历史版本1.2为550元，仅用于2025年1月1日至2026年6月30日的历史事项。", ["550元", "1.2"], "08", "A类城市|550元|750元", "历史版本", "medium")
    add("批准外发脱敏日志后，发送通道和权限有效期有什么要求？", "先获数据所有者与信息安全负责人书面批准；仅经批准的加密协作空间发送，下载权限72小时后自动失效，不得擅自改用普通邮件附件。", ["书面批准", "加密协作空间", "72小时", "普通邮件"], "09", "接收方下载权限必须在72小时后自动失效", "跨页条款", "hard")
    add("ASSET-026是什么系统，复核周期是多久？", "证书管理，每月复核，归口客户服务部。", ["证书管理", "每月", "客户服务部"], "10", "ASSET-026\n证书管理", "长表跨页", "medium")
    add("生产回退手册中，核心API成功率出现什么情形要决定回退？", "连续10分钟低于99.0%，或出现数据写入错误时，由发布负责人决定是否回退。", ["10分钟", "99.0%", "数据写入错误"], "11", "核心API成功率连续10分钟低于99.0%", "双栏阅读", "medium")
    add("XH-MNT-2026-017的维护时间和影响范围？", "2026年9月20日02:00至03:00北京时间，只影响生产报表导出；登录、审批和开放API不受影响。", ["9月20日", "02:00", "03:00", "报表导出"], "12", "维护窗口：2026年9月20日02:00至03:00（北京时间）", "OCR扫描件", "hard")
    add("海岚项目最终验收是在什么时候签署的？", "2026年9月29日。", ["2026年9月29日"], "13", "客户于2026年9月29日签署最终验收", "状态更新", "medium")
    add("海岚项目遗留缺陷DEF-026-08何时计划修复，是否已修复？", "计划2026年10月9日修复，验收纪要未证明已经修复。", ["2026年10月9日", "计划"], "13", "不应将计划修复日期当作已经修复", "时间状态", "hard")
    add("协作云2.6.0发布后API主版本是否变成/v2？", "没有，仍使用/v1；产品版本与API主版本不同。", ["/v1", "不同"], "14", "开放API仍使用/v1主版本路径", "版本区分", "medium")
    add("RATE-001对应什么HTTP状态，怎么办？", "HTTP 429，遵循Retry-After并指数退避，不能用多个应用绕过限流。", ["429", "Retry-After"], "15", "RATE-001|429|超过调用额度|按照Retry-After退避", "错误代码")
    add("P1首次响应默认多久，等于解决时间吗？", "15分钟；首次响应不等于解决时间，订单专属目标优先。", ["15分钟", "不等于解决时间"], "16", "首次响应不等于解决时间", "FAQ")
    add("部署手册中PostgreSQL备份保留期及RPO/RTO目标？", "35日；RPO 15分钟，RTO 2小时，是内部运维目标。", ["35日", "15分钟", "2小时"], "17", "PostgreSQL|每日全量与连续日志归档|35日|RPO 15分钟，RTO 2小时", "Markdown表格")
    add("CR-026-03增加了什么工作量，验收计划如何调整？", "新增跨部门审批代理，4人日；从9月25日调整到9月29日，双方授权代表书面批准。", ["4人日", "9月25日", "9月29日"], "18", "评估工作量为4人日", "变更关联", "medium")
    add("INC-2026-0912的首次响应与恢复各用了多久？", "首次响应17分钟，恢复43分钟。", ["17分钟", "43分钟"], "19", "首次响应耗时17分钟，恢复耗时43分钟", "时间线计算", "medium")
    add("九月知识库培训的时间及主讲是谁？", "2026年9月24日15:00至16:30北京时间，主讲许澜。", ["9月24日", "15:00", "16:30", "许澜"], "20", "时间：2026年9月24日15:00至16:30（北京时间）")
    add("开放平台Idempotency-Key有效期多久？", "24小时，写操作必需，重试同一操作保持不变。", ["24小时"], "15", "同一业务操作重试时保持不变，有效期24小时", "技术参数")
    # A multi-document single-turn question has evidence from both lifecycle stages.
    cases.append({"id": "XH-031", "question": "海岚项目为何改期，最后何时验收？",
                  "reference_answer": "CR-026-03新增跨部门审批代理，评估4人日，验收由9月25日调整至9月29日，并于9月29日签署最终验收。",
                  "answerable": True, "required_facts": ["CR-026-03", "4人日", "9月29日"],
                  "expected_evidence": [
                      {"document": "18_海岚项目周报与变更记录.md", "text": "评估工作量为4人日"},
                      {"document": "13_海岚项目最终验收纪要.pdf", "text": "客户于2026年9月29日签署最终验收"}],
                  "category": "跨文档关联", "difficulty": "hard"})
    for question in ("星海科技2026年的营业收入是多少？", "海岚项目的签约金额是多少？",
                     "公司CEO的私人手机号是多少？", "公司明年的薪资涨幅已经确定是多少？"):
        cases.append({"id": f"XH-{len(cases)+1:03d}", "question": question,
                      "reference_answer": "提供的资料未包含该信息，不能确定，不应推测。",
                      "answerable": False, "required_facts": [], "expected_evidence": [],
                      "category": "不可回答", "difficulty": "medium"})
    return cases


def write_readme(output: Path, manifest: dict):
    rows = "\n".join(f"| {Path(r['file']).name} | {r['category']} | {'、'.join(r['features'])} |" for r in manifest["documents"])
    body = f"""# 星海科技有限公司知识库测试集

这是虚构企业的离线小规模功能测试集：20份文档（6 DOCX、8 PDF、4 Markdown、2 TXT），不是完整中小企业规模的知识库，也不是性能压测基准。公司、人员、客户、事件和业务数值均用于测试，不具有实际管理或法律效力。

## 如何使用

1. 只上传 `documents/` 中的20个文件。不要上传本说明、`manifest.json` 或 `evaluation/`，避免答案泄漏。
2. 建议建立独立测试知识库；不要与旧的“启明数字科技”样例混合，避免近重复文档影响结果。
3. 先用19份非扫描文档验证切分与检索，再加入图片型维护通知测试OCR。扫描件必须有中文OCR环境；缺少OCR时应明确报错，不应静默产生空知识。
4. 保留2025年归档版可测试历史与当前版本区分；只想做基础检索时可以暂不上传它。
5. `evaluation/questions.jsonl` 有{manifest['single_turn_questions']}道单轮题，使用项目现有 `scripts/evaluate_answers.py` 的题集字段，可在你明确授权入库和模型调用后评估。此任务没有运行带模型的回答评测。
6. `evaluation/conversation_cases.json` 有3组多轮案例、共9轮，供会话接口或人工测试使用；它不是现有单轮评测脚本可以直接读取的格式。

## 文件目录

| 文档 | 业务类别 | 结构测试点 |
| --- | --- | --- |
{rows}

## 来源与省量方式

8份内容复用了本项目 `sample_docs/fictional_enterprise/` 的既有模拟文档：6份重排为DOCX，API和FAQ转换为Markdown。新补12份业务文档，使用离线模板和结构化事实生成，没有调用外部大模型、下载第三方语料或自动入库。

企业名、文档编号及Webhook签名示例头统一为星海版本；原始样例文件未修改。修正了继承内容中的两处一致性问题：年休假按剩余日历天数折算；恰好50000元采购属于至少2家报价档，超过50000元才进入3家报价档。历史差旅550元与现行650元是有意保留的版本差异，不是同一时间的矛盾。

`manifest.json` 包含文件类型、来源、状态、PDF页数、字符数和SHA-256校验值。扫描件的字符数取人工源文本，`embedded_text_characters=0`；其原文只在 `evaluation/scan_ground_truth.json` 中用于检查OCR，不冒充OCR结果。

## 验收状态与限制

结构、文件数量、改名一致性、证据文本存在性、扫描件无文本层及多页特性已离线检查。PDF视觉检查、DOCX渲染和真实OCR的具体状态记录在 `evaluation/generation_checks.json`；不要将内容检查或渲染代理视为真实Word分页验收。

这个集合适合检查标题、表格、跨页、OCR、版本与多轮追问，规模不足以证明中小企业规模下的吞吐或准确率。若要扩至数百/千份，建议先依据此批测试结果扩充不同业务事实，而不是重复同一文件或机械改标题。

## 重新生成

使用已具备 python-docx、ReportLab、Pillow、pypdf 的文档运行时执行 `scripts/generate_xinghai_corpus.py --render`。脚本使用本机微软雅黑字体，不将字体文件加入文档包；Poppler用于PDF预览。再生成不会触碰项目数据库或旧样例。打包由同一脚本完成，预览图片和生成器不进入ZIP。
"""
    (output / "README.md").write_text(body, encoding="utf-8")


def package(output: Path):
    archive = output.parent / (COMPANY + "_知识库测试集.zip")
    with zipfile.ZipFile(archive, "w", compression=zipfile.ZIP_DEFLATED) as bundle:
        for path in sorted(output.rglob("*")):
            if path.is_file():
                bundle.write(path, Path(output.name) / path.relative_to(output))


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--output", type=Path, default=ROOT / "output/pdf/xinghai-tech-corpus")
    parser.add_argument("--render", action="store_true", help="Render PDFs with Poppler for manual visual review")
    args = parser.parse_args()
    build(args.output.resolve(), args.render)
